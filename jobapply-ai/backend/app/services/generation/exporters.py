"""CV exporters: ATS-friendly single-column PDF (ReportLab), DOCX (python-docx) and LaTeX source.

All three render the same structured CV dict. Text is real text (no images/tables for layout) so ATS parsers can read it.
"""
from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from app.services.generation.cv_builder import shrink

SECTION_TITLES = {
    "en": {"summary": "Summary", "skills": "Skills", "experience": "Experience", "education": "Education", "projects": "Projects",
           "certifications": "Certifications", "achievements": "Achievements", "publications": "Publications", "languages": "Languages",
           "extracurriculars": "Activities", "present": "Present"},
    "bn": {"summary": "সারসংক্ষেপ", "skills": "দক্ষতা", "experience": "অভিজ্ঞতা", "education": "শিক্ষা", "projects": "প্রকল্প",
           "certifications": "সার্টিফিকেশন", "achievements": "অর্জন", "publications": "প্রকাশনা", "languages": "ভাষা",
           "extracurriculars": "কার্যক্রম", "present": "বর্তমান"},
}
TEMPLATES = {
    "classic": {"accent": "#111111", "name_size": 18, "body": 9.5, "rule": True, "heading_caps": True},
    "modern": {"accent": "#1f4e79", "name_size": 20, "body": 9.5, "rule": True, "heading_caps": False},
    "compact": {"accent": "#222222", "name_size": 16, "body": 9, "rule": False, "heading_caps": True},
}
FONTS = {"serif": ("Times-Roman", "Times-Bold", "Times-Italic", "Times New Roman"), "sans": ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Arial")}


def _dates(item: dict, lang: str) -> str:
    start, end = item.get("start_date", ""), item.get("end_date", "")
    if item.get("is_current"):
        end = SECTION_TITLES.get(lang, SECTION_TITLES["en"])["present"]
    return " – ".join(x for x in (start, end) if x)


def _t(lang: str) -> dict:
    return SECTION_TITLES.get(lang, SECTION_TITLES["en"])


# ------------------------------------------------------------------ complex scripts (Bangla) ---------------------------------

class ComplexScriptUnavailable(Exception):
    """Raised when text needs OpenType shaping (e.g. Bangla) but WeasyPrint/Pango is not installed."""


_BENGALI = re.compile("[\u0980-\u09FF]")


def needs_shaping(*objs) -> bool:
    def walk(o):
        if isinstance(o, str):
            return bool(_BENGALI.search(o))
        if isinstance(o, dict):
            return any(walk(v) for v in o.values())
        if isinstance(o, (list, tuple)):
            return any(walk(v) for v in o)
        return False

    return any(walk(o) for o in objs)


_HTML_FONTS = {"serif": '"Noto Serif Bengali","Noto Serif","FreeSerif","Times New Roman",serif', "sans": '"Noto Sans Bengali","Noto Sans","FreeSans","Arial",sans-serif'}


def _html_pdf(body_html: str, css: str, title: str) -> tuple[bytes, int]:
    try:
        import weasyprint
    except (ImportError, OSError) as exc:  # pragma: no cover - depends on system libs
        raise ComplexScriptUnavailable("WeasyPrint is not installed") from exc
    html = f"<!doctype html><html><head><meta charset='utf-8'><title>{escape(title)}</title><style>{css}</style></head><body>{body_html}</body></html>"
    doc = weasyprint.HTML(string=html).render()
    return doc.write_pdf(), len(doc.pages)


def _cv_html(cv: dict, opts: dict, scale: float) -> tuple[str, str]:
    tpl = TEMPLATES.get(opts.get("template", "classic"), TEMPLATES["classic"])
    lang = opts.get("language", "en")
    t = _t(lang)
    fam = _HTML_FONTS.get(opts.get("font", "serif"), _HTML_FONTS["serif"])
    base = tpl["body"] * scale
    e = escape
    h = cv["header"]
    css = (f"@page{{size:A4;margin:13mm 16mm 12mm 16mm}}body{{font-family:{fam};font-size:{base}pt;line-height:1.28;color:#111}}"
           f"h1{{font-size:{tpl['name_size'] * scale}pt;margin:0;color:{tpl['accent']}}}h2{{font-size:{base + 1.5}pt;margin:6pt 0 2pt;color:{tpl['accent']};"
           f"{'border-bottom:0.6pt solid ' + tpl['accent'] + ';' if tpl['rule'] else ''}}}p{{margin:1pt 0}}ul{{margin:1pt 0 2pt 14pt;padding:0}}li{{margin:0}}.row{{font-weight:bold}}.sub{{font-style:italic}}")
    parts = [f"<h1>{e(h['name'])}</h1><p>{e('  |  '.join(x for x in [h.get('email'), h.get('phone'), h.get('location'), *h.get('links', [])] if x))}</p>"]

    def sec(key, inner):
        parts.append(f"<h2>{e(t[key])}</h2>{inner}")

    def ul(items):
        return "<ul>" + "".join(f"<li>{e(i)}</li>" for i in items) + "</ul>" if items else ""

    if cv.get("summary"):
        sec("summary", f"<p>{e(cv['summary'])}</p>")
    if cv.get("skills"):
        groups: dict[str, list[str]] = {}
        for sk in cv["skills"]:
            groups.setdefault(sk["category"], []).append(sk["name"])
        sec("skills", "".join(f"<p><b>{e(c.replace('_', ' ').title())}:</b> {e(', '.join(n))}</p>" for c, n in groups.items()))
    if cv.get("experience"):
        sec("experience", "".join(f"<p class='row'>{e(x['title'])}{' — ' + e(x['company']) if x['company'] else ''} <span style='font-weight:normal'>({e(_dates(x, lang))})</span></p>{ul(x['bullets'])}" for x in cv["experience"]))
    if cv.get("education"):
        sec("education", "".join(f"<p class='row'>{e(x['degree'])}{' — ' + e(x['institution']) if x['institution'] else ''} <span style='font-weight:normal'>({e(' – '.join(y for y in (x['start_date'], x['end_date']) if y))})</span></p>"
                                  + (f"<p>CGPA/Grade: {e(x['grade'])}</p>" if x.get("grade") else "") for x in cv["education"]))
    if cv.get("projects"):
        sec("projects", "".join(f"<p class='row'>{e(x['name'])}{' — ' + e(x['description']) if x['description'] else ''}</p>{ul(x['bullets'])}" for x in cv["projects"]))
    if cv.get("certifications"):
        sec("certifications", ul([" – ".join(y for y in (c["name"], c.get("issuer", ""), c.get("date", "")) if y) for c in cv["certifications"]]))
    for key in ("achievements", "publications", "extracurriculars"):
        if cv.get(key):
            sec(key, ul(cv[key]))
    if cv.get("languages"):
        sec("languages", f"<p>{e(', '.join(cv['languages']))}</p>")
    return "".join(parts), css


# ------------------------------------------------------------------ PDF ---------------------------------------------------------

def _pdf_bytes(cv: dict, opts: dict, scale: float) -> tuple[bytes, int]:
    tpl = TEMPLATES.get(opts.get("template", "classic"), TEMPLATES["classic"])
    reg, bold, ital, _ = FONTS.get(opts.get("font", "serif"), FONTS["serif"])
    lang = opts.get("language", "en")
    t = _t(lang)
    base = tpl["body"] * scale
    acc = colors.HexColor(tpl["accent"])
    st_name = ParagraphStyle("name", fontName=bold, fontSize=tpl["name_size"] * scale, leading=tpl["name_size"] * 1.2 * scale, textColor=acc, alignment=TA_LEFT)
    st_contact = ParagraphStyle("contact", fontName=reg, fontSize=base - 0.5, leading=(base - 0.5) * 1.3, textColor=colors.HexColor("#333333"))
    st_h = ParagraphStyle("h", fontName=bold, fontSize=base + 1.5, leading=(base + 1.5) * 1.2, textColor=acc, spaceBefore=5 * scale, spaceAfter=1.5)
    st_body = ParagraphStyle("body", fontName=reg, fontSize=base, leading=base * 1.28)
    st_bul = ParagraphStyle("bul", parent=st_body, leftIndent=0)
    st_row = ParagraphStyle("row", parent=st_body, fontName=bold)
    st_sub = ParagraphStyle("sub", parent=st_body, fontName=ital, textColor=colors.HexColor("#333333"))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=13 * mm, bottomMargin=12 * mm,
                            title=f"{cv['header']['name']} – CV", author=cv["header"]["name"])
    E: list = []

    def heading(key: str):
        title = t[key].upper() if tpl["heading_caps"] and lang == "en" else t[key]
        E.append(Paragraph(escape(title), st_h))
        if tpl["rule"]:
            E.append(HRFlowable(width="100%", thickness=0.6, color=acc, spaceAfter=2))

    def bullets(items: list[str]):
        if items:
            E.append(ListFlowable([ListItem(Paragraph(escape(b), st_bul), leftIndent=10) for b in items], bulletType="bullet", start="•",
                                  leftIndent=10, bulletFontSize=base))

    h = cv["header"]
    E.append(Paragraph(escape(h["name"]), st_name))
    contact = [x for x in [h.get("email"), h.get("phone"), h.get("location"), *h.get("links", [])] if x]
    E.append(Paragraph(escape("  |  ".join(contact)), st_contact))
    if cv.get("summary"):
        heading("summary")
        E.append(Paragraph(escape(cv["summary"]), st_body))
    if cv.get("skills"):
        heading("skills")
        groups: dict[str, list[str]] = {}
        label = {"language": "Languages", "framework": "Frameworks", "ai_ml": "AI / ML & Data", "database": "Databases", "cloud": "Cloud & DevOps", "tool": "Tools",
                 "technical": "Technical", "soft": "Soft skills"}
        for s in cv["skills"]:
            groups.setdefault(s["category"], []).append(s["name"])
        for cat, names in groups.items():
            E.append(Paragraph(f"<b>{escape(label.get(cat, cat.title()))}:</b> {escape(', '.join(names))}", st_body))
    if cv.get("experience"):
        heading("experience")
        for e in cv["experience"]:
            E.append(Paragraph(f"{escape(e['title'])}" + (f" — {escape(e['company'])}" if e["company"] else "") +
                               (f"  <font name='{reg}' size='{base - 0.5}'>({escape(_dates(e, lang))})</font>" if _dates(e, lang) else ""), st_row))
            if e.get("location"):
                E.append(Paragraph(escape(e["location"]), st_sub))
            bullets(e["bullets"])
            E.append(Spacer(1, 2))
    if cv.get("education"):
        heading("education")
        for e in cv["education"]:
            line = f"{escape(e['degree'])}" + (f" — {escape(e['institution'])}" if e["institution"] else "")
            d = " – ".join(x for x in (e["start_date"], e["end_date"]) if x)
            E.append(Paragraph(line + (f"  <font name='{reg}' size='{base - 0.5}'>({escape(d)})</font>" if d else ""), st_row))
            if e.get("grade"):
                E.append(Paragraph(f"CGPA/Grade: {escape(e['grade'])}", st_body))
            bullets(e.get("details", []))
    if cv.get("projects"):
        heading("projects")
        for p in cv["projects"]:
            head = f"{escape(p['name'])}" + (f" — {escape(p['description'])}" if p["description"] else "")
            E.append(Paragraph(head, st_row))
            if p.get("technologies"):
                E.append(Paragraph("Tech: " + escape(", ".join(p["technologies"][:8])), st_sub))
            bullets(p["bullets"])
            E.append(Spacer(1, 2))
    if cv.get("certifications"):
        heading("certifications")
        bullets([" – ".join(x for x in (c["name"], c.get("issuer", ""), c.get("date", "")) if x) for c in cv["certifications"]])
    for key in ("achievements", "publications", "extracurriculars"):
        if cv.get(key):
            heading(key)
            bullets(cv[key])
    if cv.get("languages"):
        heading("languages")
        E.append(Paragraph(escape(", ".join(cv["languages"])), st_body))
    doc.build(E)
    data = buf.getvalue()
    from pypdf import PdfReader

    return data, len(PdfReader(io.BytesIO(data)).pages)


def render_pdf(cv: dict, options: dict | None = None) -> tuple[bytes, int, dict]:
    """Render and enforce the page limit by progressively trimming content. Returns (pdf, pages, final_cv)."""
    opts = options or {}
    limit = int(opts.get("pages", 1))
    current = cv
    shaped = opts.get("language", "en") != "en" or needs_shaping(cv)
    for level in range(0, 5):
        current = cv if level == 0 else shrink(cv, level)
        for scale in ((1.0, 0.95, 0.9) if level >= 2 else (1.0,)):
            if shaped:
                body, css = _cv_html(current, opts, scale)
                data, pages = _html_pdf(body, css, f"{cv['header']['name']} – CV")
            else:
                data, pages = _pdf_bytes(current, opts, scale)
            if pages <= limit:
                return data, pages, current
    return data, pages, current  # could not fit; QC will flag it


# ------------------------------------------------------------------ DOCX --------------------------------------------------------

def render_docx(cv: dict, options: dict | None = None) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    opts = options or {}
    tpl = TEMPLATES.get(opts.get("template", "classic"), TEMPLATES["classic"])
    font = FONTS.get(opts.get("font", "serif"), FONTS["serif"])[3]
    lang = opts.get("language", "en")
    t = _t(lang)
    d = docx.Document()
    for s in d.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = Mm(16)
        s.top_margin, s.bottom_margin = Mm(13), Mm(12)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(tpl["body"])
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    normal.paragraph_format.space_after = Pt(1)
    acc = RGBColor.from_string(tpl["accent"].lstrip("#").upper())

    def para(text="", bold=False, size=None, color=None, italic=False, style=None):
        p = d.add_paragraph(style=style)
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        return p

    def heading(key: str):
        p = para(t[key].upper() if tpl["heading_caps"] and lang == "en" else t[key], bold=True, size=tpl["body"] + 1.5, color=acc)
        p.paragraph_format.space_before = Pt(6)
        if tpl["rule"]:
            pPr = p._p.get_or_add_pPr()
            b = OxmlElement("w:pBdr")
            bt = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", tpl["accent"].lstrip("#"))):
                bt.set(qn(k), v)
            b.append(bt)
            pPr.append(b)

    def bullets(items):
        for b in items:
            p = d.add_paragraph(b, style="List Bullet")
            p.paragraph_format.space_after = Pt(0)

    h = cv["header"]
    para(h["name"], bold=True, size=tpl["name_size"], color=acc)
    para("  |  ".join(x for x in [h.get("email"), h.get("phone"), h.get("location"), *h.get("links", [])] if x), size=tpl["body"] - 0.5)
    if cv.get("summary"):
        heading("summary")
        para(cv["summary"])
    if cv.get("skills"):
        heading("skills")
        groups: dict[str, list[str]] = {}
        for s in cv["skills"]:
            groups.setdefault(s["category"], []).append(s["name"])
        label = {"language": "Languages", "framework": "Frameworks", "ai_ml": "AI / ML & Data", "database": "Databases", "cloud": "Cloud & DevOps", "tool": "Tools", "technical": "Technical", "soft": "Soft skills"}
        for cat, names in groups.items():
            p = d.add_paragraph()
            p.add_run(label.get(cat, cat.title()) + ": ").bold = True
            p.add_run(", ".join(names))
    if cv.get("experience"):
        heading("experience")
        for e in cv["experience"]:
            p = d.add_paragraph()
            p.add_run(e["title"] + (f" — {e['company']}" if e["company"] else "")).bold = True
            if _dates(e, lang):
                p.add_run(f"  ({_dates(e, lang)})")
            if e.get("location"):
                para(e["location"], italic=True)
            bullets(e["bullets"])
    if cv.get("education"):
        heading("education")
        for e in cv["education"]:
            p = d.add_paragraph()
            p.add_run(e["degree"] + (f" — {e['institution']}" if e["institution"] else "")).bold = True
            dt = " – ".join(x for x in (e["start_date"], e["end_date"]) if x)
            if dt:
                p.add_run(f"  ({dt})")
            if e.get("grade"):
                para(f"CGPA/Grade: {e['grade']}")
            bullets(e.get("details", []))
    if cv.get("projects"):
        heading("projects")
        for pr in cv["projects"]:
            p = d.add_paragraph()
            p.add_run(pr["name"] + (f" — {pr['description']}" if pr["description"] else "")).bold = True
            if pr.get("technologies"):
                para("Tech: " + ", ".join(pr["technologies"][:8]), italic=True)
            bullets(pr["bullets"])
    if cv.get("certifications"):
        heading("certifications")
        bullets([" – ".join(x for x in (c["name"], c.get("issuer", ""), c.get("date", "")) if x) for c in cv["certifications"]])
    for key in ("achievements", "publications", "extracurriculars"):
        if cv.get(key):
            heading(key)
            bullets(cv[key])
    if cv.get("languages"):
        heading("languages")
        para(", ".join(cv["languages"]))
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


# ------------------------------------------------------------------ LaTeX -------------------------------------------------------

_TEX = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}


def tex_escape(s: str) -> str:
    return "".join(_TEX.get(c, c) for c in s)


def render_tex(cv: dict, options: dict | None = None) -> str:
    opts = options or {}
    lang = opts.get("language", "en")
    t = _t(lang)
    h = cv["header"]
    L = [r"\documentclass[10pt,a4paper]{article}", r"\usepackage[margin=16mm]{geometry}", r"\usepackage[T1]{fontenc}", r"\usepackage{enumitem}",
         r"\usepackage[hidelinks]{hyperref}", r"\setlist[itemize]{leftmargin=*,nosep}", r"\pagestyle{empty}", r"\begin{document}",
         rf"{{\LARGE\bfseries {tex_escape(h['name'])}}}\\", tex_escape("  |  ".join(x for x in [h.get('email'), h.get('phone'), h.get('location'), *h.get('links', [])] if x)), ""]

    def sec(key):
        L.append(rf"\section*{{{tex_escape(t[key].upper())}}}")

    def items(xs):
        if xs:
            L.append(r"\begin{itemize}")
            L.extend(rf"\item {tex_escape(x)}" for x in xs)
            L.append(r"\end{itemize}")

    if cv.get("summary"):
        sec("summary"); L.append(tex_escape(cv["summary"]))
    if cv.get("skills"):
        sec("skills"); L.append(tex_escape(", ".join(s["name"] for s in cv["skills"])))
    if cv.get("experience"):
        sec("experience")
        for e in cv["experience"]:
            L.append(rf"\textbf{{{tex_escape(e['title'])}}}" + (f" --- {tex_escape(e['company'])}" if e["company"] else "") + (f" \\hfill {tex_escape(_dates(e, lang))}" if _dates(e, lang) else "") + r"\\")
            items(e["bullets"])
    if cv.get("education"):
        sec("education")
        for e in cv["education"]:
            dt = " -- ".join(x for x in (e["start_date"], e["end_date"]) if x)
            L.append(rf"\textbf{{{tex_escape(e['degree'])}}}" + (f" --- {tex_escape(e['institution'])}" if e["institution"] else "") + (f" \\hfill {tex_escape(dt)}" if dt else "") + r"\\")
            if e.get("grade"):
                L.append(tex_escape("CGPA/Grade: " + e["grade"]) + r"\\")
    if cv.get("projects"):
        sec("projects")
        for p in cv["projects"]:
            L.append(rf"\textbf{{{tex_escape(p['name'])}}}" + (f" --- {tex_escape(p['description'])}" if p["description"] else "") + r"\\")
            items(p["bullets"])
    if cv.get("certifications"):
        sec("certifications"); items([" - ".join(x for x in (c["name"], c.get("issuer", ""), c.get("date", "")) if x) for c in cv["certifications"]])
    for key in ("achievements", "publications", "extracurriculars"):
        if cv.get(key):
            sec(key); items(cv[key])
    if cv.get("languages"):
        sec("languages"); L.append(tex_escape(", ".join(cv["languages"])))
    L.append(r"\end{document}")
    return "\n".join(L)


def render_cover_letter_pdf(letter: dict, header: dict, options: dict | None = None) -> bytes:
    opts = options or {}
    if needs_shaping(letter, header) or opts.get("language", "en") != "en":
        e = escape
        fam = _HTML_FONTS.get(opts.get("font", "serif"), _HTML_FONTS["serif"])
        body = (f"<p><b>{e(header['name'])}</b><br>{e('  |  '.join(x for x in [header.get('email'), header.get('phone'), header.get('location')] if x))}</p>"
                + (f"<p>{e(letter['greeting'])}</p>" if letter.get("greeting") else "") + "".join(f"<p>{e(p)}</p>" for p in letter.get("paragraphs", []))
                + (f"<p>{e(letter['closing']).replace(chr(10), '<br>')}</p>" if letter.get("closing") else ""))
        css = f"@page{{size:A4;margin:20mm 22mm}}body{{font-family:{fam};font-size:10.5pt;line-height:1.5}}p{{margin:0 0 8pt}}"
        return _html_pdf(body, css, f"{header['name']} – Cover Letter")[0]
    reg, bold, _, _ = FONTS.get(opts.get("font", "serif"), FONTS["serif"])
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=20 * mm, bottomMargin=18 * mm,
                            title=f"{header['name']} – Cover Letter", author=header["name"])
    body = ParagraphStyle("b", fontName=reg, fontSize=10.5, leading=14.5, spaceAfter=8)
    head = ParagraphStyle("h", fontName=bold, fontSize=12, leading=15)
    E = [Paragraph(escape(header["name"]), head),
         Paragraph(escape("  |  ".join(x for x in [header.get("email"), header.get("phone"), header.get("location")] if x)), ParagraphStyle("c", parent=body, fontSize=9, textColor=colors.HexColor("#444444"))),
         Spacer(1, 10)]
    if letter.get("greeting"):
        E.append(Paragraph(escape(letter["greeting"]), body))
    for para in letter.get("paragraphs", []):
        E.append(Paragraph(escape(para), body))
    if letter.get("closing"):
        E.append(Paragraph(escape(letter["closing"]).replace("\n", "<br/>"), body))
    doc.build(E)
    return buf.getvalue()


def render_cover_letter_docx(letter: dict, header: dict) -> bytes:
    import docx
    from docx.shared import Mm, Pt

    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Mm(22)
    d.styles["Normal"].font.name, d.styles["Normal"].font.size = "Times New Roman", Pt(11)
    p = d.add_paragraph(); p.add_run(header["name"]).bold = True
    d.add_paragraph("  |  ".join(x for x in [header.get("email"), header.get("phone"), header.get("location")] if x))
    if letter.get("greeting"):
        d.add_paragraph(letter["greeting"])
    for para in letter.get("paragraphs", []):
        d.add_paragraph(para)
    if letter.get("closing"):
        d.add_paragraph(letter["closing"])
    out = io.BytesIO(); d.save(out)
    return out.getvalue()


def extract_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
