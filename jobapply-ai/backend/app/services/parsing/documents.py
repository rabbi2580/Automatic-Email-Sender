"""Turn uploaded files into text. Magic-byte sniffing, size caps, OCR for images/scanned PDFs."""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass

from app.core.config import get_settings


class DocumentError(Exception):
    """User-facing parse failure with a reason the UI can show."""

    def __init__(self, reason: str, code: str = "unreadable"):
        super().__init__(reason)
        self.reason = reason
        self.code = code


@dataclass
class ExtractedDocument:
    text: str
    kind: str  # pdf | docx | text | image
    used_ocr: bool = False
    pages: int = 1


def sniff_kind(data: bytes, filename: str = "") -> str:
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04" and (filename.lower().endswith(".docx") or b"word/" in data[:4096] or b"[Content_Types]" in data[:4096]):
        return "docx"
    if data[:8] == b"\x89PNG\r\n\x1a\n" or data[:3] == b"\xff\xd8\xff" or data[:4] in (b"RIFF", b"GIF8") or data[:2] == b"BM":
        return "image"
    try:
        data[:4096].decode("utf-8")
        return "text"
    except UnicodeDecodeError:
        raise DocumentError("Unsupported or corrupted file type. Please upload a PDF, DOCX, image or text file.", "unsupported")


def check_size(data: bytes) -> None:
    limit = get_settings().max_upload_mb * 1024 * 1024
    if len(data) > limit:
        raise DocumentError(f"File is larger than {get_settings().max_upload_mb} MB.", "too_large")
    if not data:
        raise DocumentError("The file is empty.", "empty")


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n"))
    # Common UTF-8-as-Windows-1252 corruption from copy/paste/export tools.
    if any(mark in text for mark in ("â", "Â", "ðŸ")):
        try:
            repaired = text.encode("latin1").decode("utf-8")
            if sum(text.count(x) for x in ("â", "Â", "ðŸ")) > sum(repaired.count(x) for x in ("â", "Â", "ðŸ")):
                text = repaired
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf_text(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise DocumentError("The PDF is password-protected. Remove the password and upload again.", "encrypted")
        # pdfplumber preserves line positioning better than pypdf for columns,
        # tables and mixed Unicode fonts. Keep pypdf as a dependency-light
        # fallback because some PDFs have malformed layout metadata.
        pages = []
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                pages = [(p.extract_text(layout=True) or "") for p in pdf.pages]
        except Exception:  # noqa: BLE001
            pages = [(p.extract_text() or "") for p in reader.pages]
        if not any(p.strip() for p in pages):
            pages = [(p.extract_text() or "") for p in reader.pages]
    except DocumentError:
        raise
    except (PdfReadError, Exception) as exc:  # noqa: BLE001
        raise DocumentError(f"Could not read the PDF ({type(exc).__name__}).", "corrupt")
    return "\n".join(pages), len(pages)


def _docx_text(data: bytes) -> str:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Could not read the DOCX ({type(exc).__name__}).", "corrupt")
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def ocr_image(data: bytes) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("OCR is not available on this server.", "ocr_unavailable") from exc
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
        if img.width * img.height > 40_000_000:
            raise DocumentError("The image resolution is too large.", "too_large")
        text = pytesseract.image_to_string(img.convert("L"))
    except DocumentError:
        raise
    except pytesseract.TesseractNotFoundError as exc:  # type: ignore[attr-defined]
        raise DocumentError("OCR engine (Tesseract) is not installed on this server. Paste the job text instead.", "ocr_unavailable") from exc
    except Exception as exc:  # noqa: BLE001
        raise DocumentError(f"Could not read the image ({type(exc).__name__}).", "corrupt")
    return text


def extract_text(data: bytes, filename: str = "", min_chars: int = 40) -> ExtractedDocument:
    """Extract text from a PDF/DOCX/image/text file or raise DocumentError with a user-facing reason."""
    check_size(data)
    kind = sniff_kind(data, filename)
    used_ocr = False
    pages = 1
    if kind == "pdf":
        text, pages = _pdf_text(data)
        if len(text.strip()) < min_chars:  # likely scanned: try OCR of rendered pages
            ocr = _ocr_pdf(data)
            if ocr:
                text, used_ocr = ocr, True
    elif kind == "docx":
        text = _docx_text(data)
    elif kind == "image":
        text, used_ocr = ocr_image(data), True
    else:
        text = data.decode("utf-8", errors="replace")
    text = clean_text(text)
    if len(text) < min_chars:
        raise DocumentError(
            "Not enough readable text was found. If this is a photo or scan, the image may be too blurry or low resolution.",
            "too_little_text",
        )
    return ExtractedDocument(text=text, kind=kind, used_ocr=used_ocr, pages=pages)


def _ocr_pdf(data: bytes) -> str:
    try:
        import pypdfium2 as pdfium
        import pytesseract

        pdf = pdfium.PdfDocument(data)
        out = []
        for i in range(min(len(pdf), 4)):
            img = pdf[i].render(scale=2).to_pil().convert("L")
            out.append(pytesseract.image_to_string(img))
        return "\n".join(out)
    except Exception:  # noqa: BLE001 - OCR is best-effort
        return ""
