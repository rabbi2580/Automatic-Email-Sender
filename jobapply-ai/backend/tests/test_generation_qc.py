"""Document generation + factual-integrity guarantees + QC blocking."""
import copy
import io
from datetime import date

import pytest

from app.schemas.ai import TailorSelection
from app.services.generation import exporters
from app.services.generation.cv_builder import BuildOptions, build_tailored_cv, rewrite_is_safe, shrink
from app.services.generation.letters import deterministic_cover_letter, generate_email, validate_letter
from app.services.parsing.cv_parser import parse_cv_text
from app.services.parsing.job_parser import parse_job_text
from app.services.qc import run_qc
from tests.helpers import FIX

TODAY = date(2026, 10, 2)


@pytest.fixture()
def world():
    d = parse_cv_text((FIX / "sample_cv.txt").read_text(), TODAY).model_dump()
    prof = {**{k: d[k] for k in ("full_name", "email", "phone", "location", "headline", "summary", "links", "years_experience", "languages", "achievements", "publications", "extracurriculars")},
            "target_roles": [], "preferred_locations": [], "work_preference": "any",
            "skills": [{"id": f"s{i}", "name": s["name"], "canonical": s["name"], "category": s["category"]} for i, s in enumerate(d["skills"])],
            "experiences": [{"id": f"e{i}", **e} for i, e in enumerate(d["experiences"])], "educations": [{"id": f"d{i}", **e} for i, e in enumerate(d["educations"])],
            "projects": [{"id": f"p{i}", **e} for i, e in enumerate(d["projects"])], "certifications": [{"id": f"c{i}", **e} for i, e in enumerate(d["certifications"])]}
    j = parse_job_text((FIX / "job1.txt").read_text(), TODAY).model_dump(mode="json")
    job = {"id": "j1", "company_name": j["company"], "job_title": j["job_title"], "required_skills": j["required_skills"], "preferred_skills": j["preferred_skills"],
           "tech_stack": j["tech_stack"], "responsibilities": j["responsibilities"], "application_email": j["application_email"], "application_url": ""}
    return prof, job


def full_qc(prof, job, cv, letter=None, email=None, pdf=None, pages=1, attachments=None, prior=None):
    letter = letter or deterministic_cover_letter(prof, job)
    em = email or {**generate_email(prof, job, "professional", "en", True, ["Python"]), "to_address": job["application_email"]}
    att = attachments if attachments is not None else [{"kind": "cv_pdf"}, {"kind": "cover_letter_pdf"}]
    return run_qc(profile=prof, job=job, cv=cv, pdf=pdf, pages=pages, pages_limit=1, letter=letter, email=em, attachments=att, include_letter=True, prior_sends=prior or [])


def test_tailored_cv_is_grounded_and_ordered_by_relevance(world):
    prof, job = world
    cv = build_tailored_cv(prof, job, BuildOptions(pages=1))
    assert cv["header"]["name"] == prof["full_name"]
    names = [s["name"] for s in cv["skills"]]
    assert names.index("Django") < names.index("Java")  # job-relevant skills first
    assert cv["projects"][0]["name"] == "Blood-Connect"  # Django/PostgreSQL project outranks the CNN project for this job
    assert {e["company"] for e in cv["experience"]} == {"Innovative Techworks Group"}
    assert cv["experience"][0]["start_date"] == "Jun 2026" and cv["experience"][0]["is_current"] is True
    assert full_qc(prof, job, cv)["passed"]


def test_exports_are_valid_and_ats_readable(world):
    prof, job = world
    cv = build_tailored_cv(prof, job, BuildOptions(pages=1))
    for tpl in ("classic", "modern", "compact"):
        for font in ("serif", "sans"):
            pdf, pages, _ = exporters.render_pdf(cv, {"template": tpl, "font": font, "pages": 1})
            text = exporters.extract_pdf_text(pdf)
            assert pages == 1 and prof["full_name"] in text and prof["email"] in text and "Education" in text.title() or "EDUCATION" in text
    import docx
    d = docx.Document(io.BytesIO(exporters.render_docx(cv)))
    assert prof["full_name"] in "\n".join(p.text for p in d.paragraphs)
    tex = exporters.render_tex(cv)
    assert tex.startswith("\\documentclass") and "\\end{document}" in tex and "Blood-Connect" in tex


def test_tex_escaping():
    assert exporters.tex_escape(r"R&D 100% #1 _x_ {y} \z") == r"R\&D 100\% \#1 \_x\_ \{y\} \textbackslash{}z"


def test_one_page_limit_enforced_by_trimming(world):
    prof, job = world
    big = copy.deepcopy(prof)
    for i in range(8):
        big["experiences"].append({"id": f"ex{i}", "kind": "work", "company": f"Company {i}", "title": "Engineer", "location": "", "start_date": f"Jan 20{10+i}", "end_date": f"Dec 20{10+i}",
                                   "is_current": False, "bullets": [f"Delivered feature number {j} using Python for client {i}" for j in range(6)], "technologies": ["Python"]})
    cv = build_tailored_cv(big, job, BuildOptions(pages=2))
    _, pages2, _ = exporters.render_pdf(cv, {"pages": 2})
    pdf1, pages1, final = exporters.render_pdf(cv, {"pages": 1})
    assert pages1 <= 1 or pages1 < pages2 + 5
    assert len(final["experience"]) == len(cv["experience"])  # never drops dated work history, only bullets


def test_shrink_only_removes(world):
    prof, job = world
    cv = build_tailored_cv(prof, job, BuildOptions(pages=2))
    for lvl in range(1, 5):
        s = shrink(cv, lvl)
        assert len(s["skills"]) <= len(cv["skills"]) and all(set(b) <= set(o["bullets"]) for b, o in zip((e["bullets"] for e in s["experience"]), cv["experience"]))


def test_rewrite_guard():
    assert rewrite_is_safe("Built APIs with Python", "Developed APIs using Python")
    assert not rewrite_is_safe("Built APIs with Python", "Built APIs with Python and Kubernetes")  # new skill
    assert not rewrite_is_safe("Improved accuracy", "Improved accuracy by 40%")  # new number
    assert not rewrite_is_safe("x", "")


def test_llm_selection_cannot_inject_facts(world):
    prof, job = world
    sel = TailorSelection(summary="Expert in Kubernetes and Terraform with 10 years experience", skill_ids_order=["not-a-real-id", "s0"],
                          bullet_rewrites={"e0": ["Built a Kubernetes platform serving 1M users", "Led a team of 5"]})
    cv = build_tailored_cv(prof, job, BuildOptions(pages=1), sel)
    assert "Kubernetes" not in cv["summary"]  # unsafe summary rejected → deterministic summary
    assert all("Kubernetes" not in b for e in cv["experience"] for b in e["bullets"])
    assert full_qc(prof, job, cv)["passed"]


# ---------------- QC must BLOCK tampered or wrong content ----------------

def blocked(report, check_id):
    return any(c["id"] == check_id for c in report["blocking"])


def test_qc_blocks_fabricated_skill_employer_and_dates(world):
    prof, job = world
    cv = build_tailored_cv(prof, job, BuildOptions(pages=1))

    bad = copy.deepcopy(cv); bad["skills"].append({"id": "zz", "name": "Kubernetes", "category": "cloud"})
    assert blocked(full_qc(prof, job, bad), "cv.skills")

    bad = copy.deepcopy(cv); bad["experience"][0]["company"] = "Google"
    r = full_qc(prof, job, bad); assert blocked(r, "cv.experience") and r["passed"] is False and r["summary"].startswith("Application blocked")

    bad = copy.deepcopy(cv); bad["experience"][0]["start_date"] = "Jan 2020"
    assert blocked(full_qc(prof, job, bad), "cv.experience")

    bad = copy.deepcopy(cv); bad["education"][0]["degree"] = "PhD in Computer Science"
    assert blocked(full_qc(prof, job, bad), "cv.education")

    bad = copy.deepcopy(cv); bad["experience"][0]["bullets"].append("Led a team of 12 engineers to ship a payments platform")
    assert blocked(full_qc(prof, job, bad), "cv.experience")

    bad = copy.deepcopy(cv); bad["certifications"].append({"name": "AWS Solutions Architect Professional", "issuer": "AWS", "date": "2025"})
    assert blocked(full_qc(prof, job, bad), "cv.certifications")

    bad = copy.deepcopy(cv); bad["header"]["name"] = "Someone Else"
    assert blocked(full_qc(prof, job, bad), "cv.name")

    bad = copy.deepcopy(cv); bad["meta"]["target_company"] = "Wrong Company Ltd"
    r = full_qc(prof, job, bad); assert blocked(r, "cv.target_company") and "mismatch" in r["blocking"][0]["message"].lower()


def test_qc_blocks_bad_cover_letter(world):
    prof, job = world
    cv = build_tailored_cv(prof, job, BuildOptions(pages=1))
    ok = deterministic_cover_letter(prof, job)
    assert full_qc(prof, job, cv, letter=ok)["passed"]

    for mutate, check in [
        (lambda l: l["paragraphs"].__setitem__(0, "Dear [Hiring Manager], I am excited to apply to [Company Name]."), "letter.content"),
        (lambda l: l["paragraphs"].__setitem__(0, "I am applying for a role at Unrelated Corp Ltd."), "letter.employer_refs"),
        (lambda l: l["paragraphs"].append("I have 7 years of experience building Kubernetes clusters."), "letter.content"),
        (lambda l: l["paragraphs"].append("While at Globex Corporation Inc I shipped many things for XYZ Technologies Ltd."), "letter.employer_refs"),
    ]:
        l = copy.deepcopy(ok); mutate(l)
        assert blocked(full_qc(prof, job, cv, letter=l), check), check


def test_qc_blocks_bad_email(world):
    prof, job = world
    cv = build_tailored_cv(prof, job, BuildOptions(pages=1))
    good = {**generate_email(prof, job, "professional", "en", True, ["Python"]), "to_address": job["application_email"]}
    assert full_qc(prof, job, cv, email=good)["passed"]
    assert blocked(full_qc(prof, job, cv, email={**good, "to_address": "not-an-email"}), "email.recipient_valid")
    assert blocked(full_qc(prof, job, cv, email={**good, "subject": "Hello"}), "email.subject")
    assert blocked(full_qc(prof, job, cv, email={**good, "body": good["body"] + "\n[Your Name]"}), "email.placeholders")
    assert blocked(full_qc(prof, job, cv, attachments=[{"kind": "cover_letter_pdf"}]), "email.attachment_cv")
    assert blocked(full_qc(prof, job, cv, attachments=[{"kind": "cv_pdf"}]), "email.attachment_letter")
    assert blocked(full_qc(prof, job, cv, prior=[{"to": job["application_email"], "job_title": job["job_title"]}]), "email.duplicate")
    r = full_qc(prof, job, cv, email={**good, "to_address": "other@else.com"})
    assert r["passed"] and any(w["id"] == "email.recipient_matches_post" for w in r["warnings"])  # differing recipient is a warning, not a block


def test_letter_tones_and_languages(world):
    prof, job = world
    texts = {t: "\n".join(deterministic_cover_letter(prof, job, t)["paragraphs"]) for t in ("professional", "concise", "technical", "research", "startup")}
    assert len(set(texts.values())) >= 4
    assert len(texts["concise"]) < len(texts["professional"])
    for t in texts:
        assert validate_letter(deterministic_cover_letter(prof, job, t), prof, job) == []
    bn = deterministic_cover_letter(prof, job, "professional", "bn")
    assert bn["greeting"].startswith("প্রিয়") and job["company_name"] in "\n".join(bn["paragraphs"])
    hdr = {"name": prof["full_name"], "email": prof["email"], "phone": prof["phone"], "location": prof["location"]}
    assert exporters.render_cover_letter_pdf(deterministic_cover_letter(prof, job), hdr)[:5] == b"%PDF-"
    bn_pdf = exporters.render_cover_letter_pdf(bn, hdr, {"language": "bn"})  # shaped via WeasyPrint
    assert bn_pdf[:5] == b"%PDF-"
    cv_bn, = [build_tailored_cv(prof, job, BuildOptions(pages=1, language="bn"))]
    pdf, pages, _ = exporters.render_pdf(cv_bn, {"language": "bn", "pages": 1})
    assert pdf[:5] == b"%PDF-" and pages == 1


def test_email_generation(world):
    prof, job = world
    e = generate_email(prof, job, "professional", "en", True, ["Python", "Django"])
    assert e["subject"] == "Application for Junior Software Engineer – Tarif Ul Haider Rabbi"
    assert "cover letter" in e["body"] and "XYZ Technologies Ltd" in e["body"]
    assert "cover letter" not in generate_email(prof, job, "professional", "en", False, [])["body"]


def test_qc_blocks_unknown_links_and_addresses():
    from app.services.qc import QCResult, check_links

    profile = {"email": "me@example.com", "links": {"linkedin": "https://linkedin.com/in/me", "other": []}, "projects": []}
    job = {"application_email": "hr@acme.com", "application_url": "https://acme.com/jobs/1"}
    ok = QCResult()
    check_links(ok, "Reach me at me@example.com or https://www.linkedin.com/in/me. Apply via https://acme.com/jobs/1 / hr@acme.com", profile, job, "email")
    assert ok.checks[-1]["ok"]
    bad = QCResult()
    check_links(bad, "Please visit http://evil.example/steal and mail attacker@evil.example", profile, job, "email")
    assert not bad.checks[-1]["ok"] and "evil.example" in bad.checks[-1]["message"]
